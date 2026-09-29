import io
import os
import re
from datetime import datetime, timezone
from typing import Any, Dict, List, Optional, Tuple

import requests
from dotenv import load_dotenv
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from pypdf import PdfReader
from docx import Document
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

load_dotenv()

APP_NAME = "Real-Time Resume Job Matcher"
APP_VERSION = "5.0.0"

ADZUNA_APP_ID = os.getenv("ADZUNA_APP_ID", "").strip()
ADZUNA_APP_KEY = os.getenv("ADZUNA_APP_KEY", "").strip()
ADZUNA_BASE_URL = "https://api.adzuna.com/v1/api/jobs/in/search"

DEFAULT_RESULTS = 10
MAX_RESULTS = 20
DEFAULT_MAX_JOB_AGE_DAYS = 30

app = FastAPI(
    title=APP_NAME,
    version=APP_VERSION,
    description="Real-time resume analysis and job matching using Adzuna.",
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

SKILLS = [
    "python", "java", "c", "c++", "c#", "javascript", "typescript",
    "html", "css", "react", "angular", "vue", "node.js", "express",
    "fastapi", "flask", "spring", "spring boot", "django", "sql",
    "mysql", "postgresql", "mongodb", "oracle", "git", "github",
    "docker", "kubernetes", "aws", "azure", "gcp", "machine learning",
    "deep learning", "artificial intelligence", "data science", "pandas",
    "numpy", "tensorflow", "pytorch", "scikit-learn", "rest api",
    "graphql", "iot", "linux", "jenkins", "redis", "kafka", "spark",
    "hadoop", "figma", "power bi", "tableau",
]

SKILL_ALIASES = {
    "js": "javascript",
    "ts": "typescript",
    "reactjs": "react",
    "react.js": "react",
    "node": "node.js",
    "nodejs": "node.js",
    "expressjs": "express",
    "postgres": "postgresql",
    "mongo": "mongodb",
    "ml": "machine learning",
    "ai": "artificial intelligence",
    "sklearn": "scikit-learn",
    "scikit learn": "scikit-learn",
    "restful api": "rest api",
    "restful apis": "rest api",
    "springboot": "spring boot",
}


def normalize_text(text: str) -> str:
    text = text or ""
    replacements = {
        "\u2013": "-",
        "\u2014": "-",
        "\u2018": "'",
        "\u2019": "'",
        "\u201c": '"',
        "\u201d": '"',
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    return re.sub(r"\s+", " ", text).strip().lower()


def normalize_skill(skill: str) -> str:
    skill = normalize_text(skill)
    return SKILL_ALIASES.get(skill, skill)


def contains_skill(text: str, skill: str) -> bool:
    text = normalize_text(text)
    skill = normalize_skill(skill)
    escaped = re.escape(skill)
    pattern = rf"(?<![a-z0-9]){escaped}(?![a-z0-9])"
    return re.search(pattern, text, flags=re.IGNORECASE) is not None


def extract_skills(text: str) -> List[str]:
    found = []
    for skill in SKILLS:
        if contains_skill(text, skill):
            found.append(normalize_skill(skill))
    return list(dict.fromkeys(found))


SECTION_ALIASES = {
    "education": [
        "education",
        "academic background",
        "academic qualifications",
    ],
    "experience": [
        "experience",
        "work experience",
        "professional experience",
        "employment",
    ],
    "projects": [
        "projects",
        "academic projects",
        "personal projects",
    ],
    "certifications": [
        "certifications",
        "certificates",
        "achievements",
    ],
    "skills": [
        "skills",
        "technical skills",
        "technical skill",
    ],
}


def extract_pdf_text(file_bytes: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(file_bytes))
        return "\n".join(
            page.extract_text() or "" for page in reader.pages
        ).strip()
    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail=f"Could not read PDF: {exc}",
        )


def extract_docx_text(file_bytes: bytes) -> str:
    try:
        document = Document(io.BytesIO(file_bytes))
        parts = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        for table in document.tables:
            for row in table.rows:
                row_values = [
                    cell.text.strip()
                    for cell in row.cells
                    if cell.text.strip()
                ]
                if row_values:
                    parts.append(" | ".join(row_values))

        return "\n".join(parts).strip()
    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail=f"Could not read DOCX: {exc}",
        )


def extract_sections(text: str) -> Dict[str, str]:
    lines = [
        line.strip()
        for line in text.splitlines()
        if line.strip()
    ]

    sections = {
        key: []
        for key in SECTION_ALIASES
    }

    current_section = None

    for line in lines:
        clean = normalize_text(line).strip(":- ")

        detected = next(
            (
                name
                for name, aliases in SECTION_ALIASES.items()
                if clean in aliases
            ),
            None,
        )

        if detected:
            current_section = detected
            continue

        if current_section:
            sections[current_section].append(line)

    return {
        key: "\n".join(value).strip()
        for key, value in sections.items()
    }


def analyze_resume(text: str) -> Dict[str, Any]:
    return {
        "characters": len(text),
        "skills": extract_skills(text),
        "sections": extract_sections(text),
    }


async def read_resume_file(file: UploadFile) -> str:
    filename = (file.filename or "").lower()

    if not filename.endswith((".pdf", ".docx")):
        raise HTTPException(
            status_code=400,
            detail="Only PDF and DOCX resumes are supported.",
        )

    file_bytes = await file.read()

    if not file_bytes:
        raise HTTPException(
            status_code=400,
            detail="Uploaded resume is empty.",
        )

    if filename.endswith(".pdf"):
        text = extract_pdf_text(file_bytes)
    else:
        text = extract_docx_text(file_bytes)

    if len(text.strip()) < 50:
        raise HTTPException(
            status_code=400,
            detail=(
                "Very little text was extracted from the resume. "
                "The PDF may be scanned and require OCR."
            ),
        )

    return text


def validate_adzuna_credentials() -> None:
    if not ADZUNA_APP_ID or not ADZUNA_APP_KEY:
        raise HTTPException(
            status_code=500,
            detail=(
                "Adzuna credentials are missing. "
                "Check backend/.env."
            ),
        )


def calculate_age_days(
    created: Optional[str],
) -> Optional[int]:
    if not created:
        return None

    try:
        created_date = datetime.fromisoformat(
            created.replace("Z", "+00:00")
        )
        age = (
            datetime.now(timezone.utc) - created_date
        ).days
        return max(age, 0)
    except (ValueError, TypeError):
        return None


def get_adzuna_jobs(
    query: str,
    location: str,
    results: int = DEFAULT_RESULTS,
    max_job_age_days: Optional[int] = DEFAULT_MAX_JOB_AGE_DAYS,
) -> List[Dict[str, Any]]:
    validate_adzuna_credentials()

    query = query.strip()
    location = location.strip()

    if not query:
        raise HTTPException(
            status_code=400,
            detail="Job search query cannot be empty.",
        )

    results = max(
        1,
        min(results, MAX_RESULTS),
    )

    url = f"{ADZUNA_BASE_URL}/1"

    params = {
        "app_id": ADZUNA_APP_ID,
        "app_key": ADZUNA_APP_KEY,
        "results_per_page": results,
        "what": query,
        "where": location,
        "content-type": "application/json",
    }

    try:
        response = requests.get(
            url,
            params=params,
            timeout=20,
        )
    except requests.RequestException as exc:
        raise HTTPException(
            status_code=502,
            detail=f"Could not connect to Adzuna: {exc}",
        )

    if response.status_code != 200:
        raise HTTPException(
            status_code=502,
            detail=(
                f"Adzuna returned HTTP "
                f"{response.status_code}: "
                f"{response.text[:500]}"
            ),
        )

    try:
        data = response.json()
    except ValueError:
        raise HTTPException(
            status_code=502,
            detail="Adzuna returned invalid JSON.",
        )

    jobs = []

    for item in data.get("results", []):
        created = item.get("created")
        age_days = calculate_age_days(created)

        if (
            max_job_age_days is not None
            and age_days is not None
            and age_days > max_job_age_days
        ):
            continue

        company = item.get("company") or {}
        location_data = item.get("location") or {}
        category = item.get("category") or {}

        jobs.append(
            {
                "id": str(item.get("id", "")),
                "title": item.get(
                    "title",
                    "Untitled job",
                ),
                "company": (
                    company.get("display_name")
                    or "Company not specified"
                ),
                "location": (
                    location_data.get("display_name")
                    or location
                    or "Location not specified"
                ),
                "description": item.get(
                    "description"
                ) or "",
                "category": (
                    category.get("label")
                    or category.get("tag")
                    or ""
                ),
                "contract_type": item.get(
                    "contract_type"
                ),
                "contract_time": item.get(
                    "contract_time"
                ),
                "salary_min": item.get(
                    "salary_min"
                ),
                "salary_max": item.get(
                    "salary_max"
                ),
                "salary_is_predicted": item.get(
                    "salary_is_predicted"
                ),
                "created": created,
                "age_days": age_days,
                "apply_url": (
                    item.get("redirect_url")
                    or item.get("application_url")
                    or ""
                ),
                "source": "Adzuna",
            }
        )

    return jobs


ROLE_FAMILIES = {
    "software": [
        "software",
        "developer",
        "engineer",
        "programmer",
        "full stack",
        "backend",
        "frontend",
        "web developer",
    ],
    "data": [
        "data",
        "analyst",
        "machine learning",
        "artificial intelligence",
        "data scientist",
    ],
    "testing": [
        "sdet",
        "qa",
        "test engineer",
        "quality assurance",
    ],
}


def role_is_relevant(
    title: str,
    query: str,
) -> bool:
    title = normalize_text(title)
    query = normalize_text(query)

    query_words = re.findall(
        r"[a-z0-9+#.]+",
        query,
    )

    if not query_words:
        return True

    if any(
        len(word) >= 2 and word in title
        for word in query_words
    ):
        return True

    for family_terms in ROLE_FAMILIES.values():
        query_matches_family = any(
            term in query
            for term in family_terms
        )
        title_matches_family = any(
            term in title
            for term in family_terms
        )

        if (
            query_matches_family
            and title_matches_family
        ):
            return True

    return False


def calculate_semantic_similarity(
    resume_text: str,
    job_title: str,
    job_description: str,
) -> Tuple[Optional[float], str]:
    resume_text = (resume_text or "").strip()

    job_text = (
        f"{job_title or ''}\n"
        f"{job_description or ''}"
    ).strip()

    if not resume_text or not job_text:
        return (
            None,
            "Semantic similarity unavailable.",
        )

    try:
        vectorizer = TfidfVectorizer(
            lowercase=True,
            ngram_range=(1, 2),
            max_features=5000,
        )

        matrix = vectorizer.fit_transform(
            [
                resume_text,
                job_text,
            ]
        )

        similarity = cosine_similarity(
            matrix[0:1],
            matrix[1:2],
        )[0][0]

        percentage = round(
            max(
                0.0,
                min(
                    1.0,
                    float(similarity),
                ),
            )
            * 100,
            2,
        )

        return (
            percentage,
            "TF-IDF + cosine similarity",
        )

    except ValueError:
        return (
            None,
            "Not enough usable text for semantic matching.",
        )


def calculate_skill_match(
    resume_skills: List[str],
    job_skills: List[str],
) -> Tuple[
    List[str],
    List[str],
    List[str],
    Optional[float],
]:
    resume_set = {
        normalize_skill(skill)
        for skill in resume_skills
    }

    job_set = {
        normalize_skill(skill)
        for skill in job_skills
    }

    matching = sorted(
        resume_set.intersection(job_set)
    )

    missing = sorted(
        job_set.difference(resume_set)
    )

    additional = sorted(
        resume_set.difference(job_set)
    )

    if not job_set:
        return (
            matching,
            missing,
            additional,
            None,
        )

    percentage = (
        len(matching)
        / len(job_set)
    ) * 100

    return (
        matching,
        missing,
        additional,
        round(percentage, 2),
    )


def calculate_title_relevance(
    resume_text: str,
    job_title: str,
) -> float:
    resume_text = normalize_text(resume_text)
    job_title = normalize_text(job_title)

    words = [
        word
        for word in re.findall(
            r"[a-z0-9+#.]+",
            job_title,
        )
        if len(word) >= 2
    ]

    if not words:
        return 0.0

    matches = sum(
        1
        for word in words
        if word in resume_text
    )

    return round(
        min(
            100.0,
            (matches / len(words)) * 100,
        ),
        2,
    )


def calculate_evidence_score(
    resume_text: str,
    job_text: str,
) -> float:
    resume_words = set(
        re.findall(
            r"[a-z0-9+#.]{3,}",
            normalize_text(resume_text),
        )
    )

    job_words = set(
        re.findall(
            r"[a-z0-9+#.]{3,}",
            normalize_text(job_text),
        )
    )

    if not resume_words or not job_words:
        return 0.0

    overlap = resume_words.intersection(
        job_words
    )

    return round(
        min(
            100.0,
            (
                len(overlap)
                / len(job_words)
            ) * 100,
        ),
        2,
    )


def build_explanation(
    matching_skills: List[str],
    missing_skills: List[str],
    semantic_score: Optional[float],
    job_skills_available: bool,
) -> str:
    if job_skills_available:
        if matching_skills and missing_skills:
            explanation = (
                "Recognized skill overlap: "
                + ", ".join(matching_skills[:6])
                + ". Potential skill gaps in available "
                "job text: "
                + ", ".join(missing_skills[:6])
                + "."
            )
        elif matching_skills:
            explanation = (
                "Recognized skill overlap: "
                + ", ".join(matching_skills[:6])
                + "."
            )
        elif missing_skills:
            explanation = (
                "No matching resume skills were "
                "recognized. Available job text "
                "mentions: "
                + ", ".join(missing_skills[:6])
                + "."
            )
        else:
            explanation = (
                "No recognizable skills were found "
                "in the available job text."
            )
    else:
        explanation = (
            "The available job snippet did not "
            "contain recognizable skills, so a "
            "reliable skill-gap comparison cannot "
            "be made."
        )

    if semantic_score is not None:
        explanation += (
            " Text similarity from the available "
            f"job text is {semantic_score}%."
        )

    return explanation


def match_resume_to_job(
    resume_text: str,
    resume_skills: List[str],
    job: Dict[str, Any],
) -> Dict[str, Any]:
    title = job.get("title", "")
    description = job.get("description", "")

    job_text = (
        f"{title}\n"
        f"{description}"
    ).strip()

    job_skills = extract_skills(
        job_text
    )

    (
        matching_skills,
        missing_skills,
        additional_resume_skills,
        skill_percentage,
    ) = calculate_skill_match(
        resume_skills,
        job_skills,
    )

    title_score = calculate_title_relevance(
        resume_text,
        title,
    )

    (
        semantic_score,
        semantic_method,
    ) = calculate_semantic_similarity(
        resume_text,
        title,
        description,
    )

    evidence_score = calculate_evidence_score(
        resume_text,
        job_text,
    )

    components = []

    if skill_percentage is not None:
        components.append(
            (
                "skill_overlap",
                skill_percentage,
                45,
            )
        )

    components.append(
        (
            "title_relevance",
            title_score,
            20,
        )
    )

    if semantic_score is not None:
        components.append(
            (
                "semantic_similarity",
                semantic_score,
                25,
            )
        )

    components.append(
        (
            "evidence_overlap",
            evidence_score,
            10,
        )
    )

    total_weight = sum(
        weight
        for _, _, weight in components
    )

    if total_weight:
        final_score = round(
            sum(
                value * weight
                for _, value, weight in components
            )
            / total_weight,
            2,
        )
    else:
        final_score = None

    description_length = len(
        description.strip()
    )

    if description_length >= 1000:
        data_quality = "medium"
    elif description_length >= 250:
        data_quality = "low"
    else:
        data_quality = "very_low"

    score_note = (
        "Score combines recognizable skill overlap, "
        "job-title relevance, TF-IDF semantic "
        "similarity, and available textual evidence. "
        "It is an estimate, not a hiring guarantee."
    )

    if description_length < 1000:
        score_note += (
            " Adzuna's search endpoint provides only "
            "a snippet of the job description, so the "
            "score does not represent the employer's "
            "complete job requirements."
        )

    return {
        "match_percentage": final_score,
        "score_note": score_note,
        "explanation": build_explanation(
            matching_skills,
            missing_skills,
            semantic_score,
            bool(job_skills),
        ),
        "data_quality": data_quality,
        "available_job_text_characters": (
            description_length
        ),
        "resume_skills": sorted(
            {
                normalize_skill(skill)
                for skill in resume_skills
            }
        ),
        "job_required_skills": sorted(
            {
                normalize_skill(skill)
                for skill in job_skills
            }
        ),
        "matching_skills": matching_skills,
        "missing_skills": missing_skills,
        "additional_resume_skills": (
            additional_resume_skills
        ),
        "semantic_similarity_percentage": (
            semantic_score
        ),
        "semantic_method": semantic_method,
        "score_breakdown": {
            "skill_overlap_percentage": (
                skill_percentage
            ),
            "skill_weight": 45,
            "title_relevance_percentage": (
                title_score
            ),
            "title_weight": 20,
            "semantic_similarity_percentage": (
                semantic_score
            ),
            "semantic_weight": 25,
            "evidence_overlap_percentage": (
                evidence_score
            ),
            "evidence_weight": 10,
            "total_score": final_score,
        },
    }


@app.get("/")
def root():
    return {
        "application": APP_NAME,
        "version": APP_VERSION,
        "status": "running",
        "features": [
            "PDF resume parsing",
            "DOCX resume parsing",
            "resume skill extraction",
            "real Adzuna job search",
            "job freshness filtering",
            "TF-IDF semantic matching",
            "explainable skill matching",
            "real application links",
        ],
    }


@app.get("/health")
def health():
    return {
        "status": "healthy",
        "adzuna_configured": bool(
            ADZUNA_APP_ID
            and ADZUNA_APP_KEY
        ),
    }


@app.post("/analyze-resume")
async def analyze_resume_endpoint(
    file: UploadFile = File(...),
):
    resume_text = await read_resume_file(
        file
    )

    analysis = analyze_resume(
        resume_text
    )

    return {
        "filename": file.filename,
        "resume_characters": (
            analysis["characters"]
        ),
        "resume_skills": (
            analysis["skills"]
        ),
        "resume_sections": (
            analysis["sections"]
        ),
    }


@app.get("/jobs/search")
def search_jobs(
    query: str,
    location: str = "",
    results: int = DEFAULT_RESULTS,
    max_job_age_days: Optional[int] = (
        DEFAULT_MAX_JOB_AGE_DAYS
    ),
):
    if (
        max_job_age_days is not None
        and max_job_age_days < 0
    ):
        raise HTTPException(
            status_code=400,
            detail=(
                "max_job_age_days cannot be negative."
            ),
        )

    jobs = get_adzuna_jobs(
        query=query,
        location=location,
        results=results,
        max_job_age_days=max_job_age_days,
    )

    return {
        "source": "Adzuna",
        "query": query,
        "location": location,
        "jobs_found": len(jobs),
        "jobs": jobs,
    }


@app.post("/match-jobs")
async def match_jobs(
    file: UploadFile = File(...),
    query: str = "software developer",
    location: str = "",
    results: int = DEFAULT_RESULTS,
    max_job_age_days: Optional[int] = (
        DEFAULT_MAX_JOB_AGE_DAYS
    ),
):
    if (
        max_job_age_days is not None
        and max_job_age_days < 0
    ):
        raise HTTPException(
            status_code=400,
            detail=(
                "max_job_age_days cannot be negative."
            ),
        )

    resume_text = await read_resume_file(
        file
    )

    resume_analysis = analyze_resume(
        resume_text
    )

    jobs = get_adzuna_jobs(
        query=query,
        location=location,
        results=results,
        max_job_age_days=max_job_age_days,
    )

    relevant_jobs = [
        job
        for job in jobs
        if role_is_relevant(
            job.get("title", ""),
            query,
        )
    ]

    jobs_to_match = (
        relevant_jobs
        if relevant_jobs
        else jobs
    )

    matched_jobs = []

    for job in jobs_to_match:
        matching = match_resume_to_job(
            resume_text=resume_text,
            resume_skills=resume_analysis[
                "skills"
            ],
            job=job,
        )

        matched_jobs.append(
            {
                **job,
                "matching": matching,
            }
        )

    matched_jobs.sort(
        key=lambda job: (
            job["matching"][
                "match_percentage"
            ]
            if job["matching"][
                "match_percentage"
            ] is not None
            else -1
        ),
        reverse=True,
    )

    return {
        "source": "Adzuna",
        "resume_filename": file.filename,
        "resume_characters": len(
            resume_text
        ),
        "resume_skills": resume_analysis[
            "skills"
        ],
        "resume_sections": (
            resume_analysis["sections"]
        ),
        "query": query,
        "location": location,
        "max_job_age_days": (
            max_job_age_days
        ),
        "jobs_found": len(
            matched_jobs
        ),
        "jobs": matched_jobs,
    }
