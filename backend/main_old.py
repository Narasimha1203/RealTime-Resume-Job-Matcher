import io
import os
import re
from datetime import datetime, timezone
from typing import Optional

import requests
from dotenv import load_dotenv
from fastapi import FastAPI, UploadFile, File, HTTPException
from pypdf import PdfReader
from docx import Document


# ============================================================
# CONFIGURATION
# ============================================================

load_dotenv()

ADZUNA_APP_ID = os.getenv("ADZUNA_APP_ID")
ADZUNA_APP_KEY = os.getenv("ADZUNA_APP_KEY")

app = FastAPI(
    title="Real-Time Resume Job Matcher",
    description="Real resume analysis and real-time job matching using Adzuna",
    version="2.0.0"
)


# ============================================================
# SKILL DATABASE
# ============================================================

SKILLS = [
    # Programming
    "python",
    "java",
    "c",
    "c++",
    "c#",
    "javascript",
    "typescript",
    "go",
    "golang",
    "ruby",
    "php",
    "kotlin",
    "swift",

    # Web
    "html",
    "css",
    "react",
    "angular",
    "vue",
    "next.js",
    "node.js",
    "express",
    "bootstrap",
    "tailwind",

    # Backend
    "fastapi",
    "flask",
    "django",
    "spring",
    "spring boot",
    "rest api",
    "restful api",
    "api development",

    # Databases
    "sql",
    "mysql",
    "postgresql",
    "mongodb",
    "oracle",
    "sqlite",
    "redis",

    # Cloud / DevOps
    "aws",
    "azure",
    "google cloud",
    "gcp",
    "docker",
    "kubernetes",
    "jenkins",
    "ci/cd",
    "terraform",

    # Data / AI
    "machine learning",
    "deep learning",
    "artificial intelligence",
    "data science",
    "data analysis",
    "pandas",
    "numpy",
    "scikit-learn",
    "tensorflow",
    "pytorch",
    "nlp",
    "computer vision",

    # Tools
    "git",
    "github",
    "gitlab",
    "jira",
    "postman",
    "linux",

    # Other
    "iot",
    "embedded systems",
    "microcontrollers",
    "cybersecurity",
    "devops",
    "agile",
    "scrum",
]


# ============================================================
# SKILL ALIASES
# ============================================================

SKILL_ALIASES = {
    "js": "javascript",
    "ts": "typescript",
    "reactjs": "react",
    "react.js": "react",
    "node": "node.js",
    "nodejs": "node.js",
    "expressjs": "express",
    "postgres": "postgresql",
    "mongo": "mongodb",
    "ml": "machine learning",
    "ai": "artificial intelligence",
    "dl": "deep learning",
    "scikit learn": "scikit-learn",
    "aws cloud": "aws",
    "google cloud platform": "google cloud",
    "gcp cloud": "gcp",
    "rest": "rest api",
}


# ============================================================
# GENERAL HELPERS
# ============================================================

def normalize_text(text: str) -> str:
    """
    Convert text to a normalized lowercase representation.
    """
    if not text:
        return ""

    text = text.lower()

    # Normalize common separators
    text = text.replace("–", "-")
    text = text.replace("—", "-")

    # Keep useful characters such as + and #
    text = re.sub(r"\s+", " ", text)

    return text.strip()

def contains_skill(text: str, skill: str) -> bool:
    """
    Check whether a skill exists as a meaningful phrase.
    """

    text = normalize_text(text)
    skill = normalize_text(skill)

    if not text or not skill:
        return False

    escaped = re.escape(skill)

    return bool(
        re.search(
            rf"(?<![a-z0-9]){escaped}(?![a-z0-9])",
            text
        )
    )

def normalize_skill(skill: str) -> str:
    skill = normalize_text(skill)
    return SKILL_ALIASES.get(skill, skill)


# ============================================================
# RESUME TEXT EXTRACTION
# ============================================================

def extract_pdf_text(file_bytes: bytes) -> str:
    """
    Extract text from a PDF resume.
    """
    try:
        reader = PdfReader(io.BytesIO(file_bytes))

        pages = []

        for page in reader.pages:
            try:
                page_text = page.extract_text() or ""
                pages.append(page_text)
            except Exception:
                continue

        return "\n".join(pages).strip()

    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail=f"Could not read PDF resume: {str(exc)}"
        )


def extract_docx_text(file_bytes: bytes) -> str:
    """
    Extract text from a DOCX resume.
    """
    try:
        document = Document(io.BytesIO(file_bytes))

        paragraphs = []

        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                paragraphs.append(paragraph.text)

        # Also inspect tables because many resumes use tables.
        for table in document.tables:
            for row in table.rows:
                row_text = []

                for cell in row.cells:
                    if cell.text.strip():
                        row_text.append(cell.text.strip())

                if row_text:
                    paragraphs.append(" ".join(row_text))

        return "\n".join(paragraphs).strip()

    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail=f"Could not read DOCX resume: {str(exc)}"
        )


def extract_resume_text(filename: str, file_bytes: bytes) -> str:
    """
    Choose the correct parser according to file extension.
    """
    filename_lower = filename.lower()

    if filename_lower.endswith(".pdf"):
        return extract_pdf_text(file_bytes)

    if filename_lower.endswith(".docx"):
        return extract_docx_text(file_bytes)

    raise HTTPException(
        status_code=400,
        detail="Only PDF and DOCX resumes are supported."
    )


# ============================================================
# SKILL EXTRACTION
# ============================================================

def extract_skills(text: str) -> list[str]:
    """
    Detect recognized skills from text.
    """
    normalized = normalize_text(text)

    found = set()

    for skill in SKILLS:
        if contains_skill(normalized, skill):
            found.add(normalize_skill(skill))

    return sorted(found)


# ============================================================
# RESUME SECTION EXTRACTION
# ============================================================

SECTION_PATTERNS = {
    "education": [
        "education",
        "academic background",
        "qualification",
        "qualifications"
    ],
    "experience": [
        "experience",
        "work experience",
        "professional experience",
        "employment"
    ],
    "projects": [
        "projects",
        "academic projects",
        "personal projects"
    ],
    "certifications": [
        "certifications",
        "certificates",
        "courses",
        "achievements"
    ],
}


def extract_section(text: str, section_names: list[str]) -> str:
    """
    Try to extract a section from a resume.
    """
    if not text:
        return ""

    lines = [line.strip() for line in text.splitlines() if line.strip()]

    start_index = None

    for index, line in enumerate(lines):
        normalized_line = normalize_text(line)

        # Remove common punctuation.
        normalized_line = re.sub(r"[:\-|]+$", "", normalized_line).strip()

        if normalized_line in section_names:
            start_index = index
            break

    if start_index is None:
        return ""

    collected = []

    for index in range(start_index + 1, len(lines)):
        current = normalize_text(lines[index])

        is_new_section = False

        for other_names in SECTION_PATTERNS.values():
            for name in other_names:
                if current == normalize_text(name):
                    is_new_section = True
                    break

            if is_new_section:
                break

        if is_new_section:
            break

        collected.append(lines[index])

    return "\n".join(collected).strip()


def extract_resume_sections(text: str) -> dict:
    sections = {}

    for section_name, patterns in SECTION_PATTERNS.items():
        normalized_patterns = [normalize_text(p) for p in patterns]

        sections[section_name] = extract_section(
            text,
            normalized_patterns
        )

    return sections


# ============================================================
# RESUME ANALYSIS
# ============================================================

def analyze_resume_text(text: str, filename: str) -> dict:
    skills = extract_skills(text)

    sections = extract_resume_sections(text)

    return {
        "resume_filename": filename,
        "resume_characters": len(text),
        "resume_skills": skills,
        "resume_sections": sections,
    }


# ============================================================
# API: RESUME ANALYSIS
# ============================================================

@app.post("/analyze-resume")
async def analyze_resume(file: UploadFile = File(...)):
    """
    Analyze a PDF/DOCX resume.
    """
    if not file.filename:
        raise HTTPException(
            status_code=400,
            detail="No filename was provided."
        )

    file_bytes = await file.read()

    if not file_bytes:
        raise HTTPException(
            status_code=400,
            detail="The uploaded file is empty."
        )

    resume_text = extract_resume_text(
        file.filename,
        file_bytes
    )

    if len(resume_text.strip()) < 20:
        raise HTTPException(
            status_code=400,
            detail=(
                "Very little text could be extracted from this resume. "
                "If it is a scanned/image-only PDF, OCR will be needed."
            )
        )

    result = analyze_resume_text(
        resume_text,
        file.filename
    )

    return {
        "status": "success",
        **result
    }


# ============================================================
# ADZUNA CONFIGURATION
# ============================================================

def validate_adzuna_credentials():
    if not ADZUNA_APP_ID or not ADZUNA_APP_KEY:
        raise HTTPException(
            status_code=500,
            detail=(
                "Adzuna credentials are missing. "
                "Check backend/.env"
            )
        )


def calculate_job_age_days(created: Optional[str]) -> Optional[int]:
    """
    Calculate job age using Adzuna's created timestamp.
    """
    if not created:
        return None

    try:
        created_dt = datetime.fromisoformat(
            created.replace("Z", "+00:00")
        )

        now = datetime.now(timezone.utc)

        age_seconds = (
            now - created_dt
        ).total_seconds()

        age_days = int(max(0, age_seconds // 86400))

        return age_days

    except Exception:
        return None


# ============================================================
# JOB TITLE RELEVANCE
# ============================================================

def title_matches_query(
    job_title: str,
    query: str
) -> bool:
    """
    Check whether a job title is relevant to the requested role.

    Allows common variations such as:
    Software Developer
    Software Engineer
    Application Developer
    Full Stack Developer
    Web Developer
    Backend Developer
    Frontend Developer
    """

    title = normalize_text(job_title)
    search = normalize_text(query)

    if not title or not search:
        return False

    # --------------------------------------------------------
    # Leadership roles
    # --------------------------------------------------------
    leadership_terms = {
        "director",
        "senior director",
        "vice president",
        "vp",
        "head",
        "chief",
        "manager",
        "management",
        "executive"
    }

    requested_leadership = any(
        term in search
        for term in leadership_terms
    )

    if not requested_leadership:
        for term in leadership_terms:
            if term in title:
                return False

    # --------------------------------------------------------
    # Common software-development role families
    # --------------------------------------------------------
    role_families = {
        "software developer": [
            "software developer",
            "software engineer",
            "application developer",
            "application engineer",
            "developer",
            "programmer"
        ],

        "software engineer": [
            "software engineer",
            "software developer",
            "application engineer",
            "application developer"
        ],

        "web developer": [
            "web developer",
            "frontend developer",
            "front end developer",
            "backend developer",
            "back end developer",
            "full stack developer",
            "full-stack developer"
        ],

        "frontend developer": [
            "frontend developer",
            "front end developer",
            "ui developer",
            "web developer"
        ],

        "backend developer": [
            "backend developer",
            "back end developer",
            "server side developer",
            "software engineer"
        ],

        "full stack developer": [
            "full stack developer",
            "full-stack developer",
            "software engineer",
            "web developer"
        ],

        "python developer": [
            "python developer",
            "python engineer",
            "software engineer"
        ],

        "java developer": [
            "java developer",
            "java engineer",
            "software engineer"
        ]
    }

    # --------------------------------------------------------
    # Check role family
    # --------------------------------------------------------
    for family, variations in role_families.items():

        if family in search:

            for variation in variations:

                if variation in title:
                    return True

            return False

    # --------------------------------------------------------
    # Generic fallback
    # --------------------------------------------------------
    query_words = set(
        word
        for word in re.findall(
            r"[a-z0-9+#.]+",
            search
        )
        if len(word) >= 3
    )

    if not query_words:
        return True

    return any(
        word in title
        for word in query_words
    )


# ============================================================
# ADZUNA JOB SEARCH
# ============================================================

def fetch_adzuna_jobs(
    query: str,
    location: str,
    page: int = 1,
    results_per_page: int = 10,
    max_age_days: int = 30
) -> list[dict]:

    validate_adzuna_credentials()

    if not query.strip():
        raise HTTPException(
            status_code=400,
            detail="Job query cannot be empty."
        )

    if results_per_page < 1 or results_per_page > 50:
        raise HTTPException(
            status_code=400,
            detail="results_per_page must be between 1 and 50."
        )

    if page < 1:
        raise HTTPException(
            status_code=400,
            detail="page must be at least 1."
        )

    # --------------------------------------------------------
    # Adzuna API URL
    # --------------------------------------------------------
    url = (
        "https://api.adzuna.com/v1/api/jobs/in/search/"
        f"{page}"
    )

    # --------------------------------------------------------
    # API parameters
    # --------------------------------------------------------
    params = {
        "app_id": ADZUNA_APP_ID,
        "app_key": ADZUNA_APP_KEY,
        "results_per_page": results_per_page,
        "what": query.strip(),
        "content-type": "application/json",
        "sort_by": "date"
    }

    if location and location.strip():
        params["where"] = location.strip()

    # --------------------------------------------------------
    # Call Adzuna
    # --------------------------------------------------------
    try:

        response = requests.get(
            url,
            params=params,
            timeout=20
        )

    except requests.RequestException as exc:

        raise HTTPException(
            status_code=502,
            detail=(
                f"Could not connect to Adzuna: {str(exc)}"
            )
        )

    # --------------------------------------------------------
    # Check API response
    # --------------------------------------------------------
    if response.status_code != 200:

        raise HTTPException(
            status_code=502,
            detail=(
                f"Adzuna returned HTTP "
                f"{response.status_code}: "
                f"{response.text[:500]}"
            )
        )

    # --------------------------------------------------------
    # Parse JSON
    # --------------------------------------------------------
    try:

        data = response.json()

    except ValueError:

        raise HTTPException(
            status_code=502,
            detail="Adzuna returned an invalid JSON response."
        )

    jobs = []

    # --------------------------------------------------------
    # Process jobs
    # --------------------------------------------------------
    for job in data.get("results", []):

        title = job.get("title") or ""
        created = job.get("created")

        age_days = calculate_job_age_days(
            created
        )

        # ----------------------------------------------------
        # Freshness filter
        # ----------------------------------------------------
        if (
            age_days is not None
            and age_days > max_age_days
        ):
            continue

        # ----------------------------------------------------
        # Relevance filter
        # ----------------------------------------------------
        if not title_matches_query(
            title,
            query
        ):
            continue

        # ----------------------------------------------------
        # Job description
        # ----------------------------------------------------
        description = (
            job.get("description")
            or ""
        )

        # ----------------------------------------------------
        # Application URL
        # ----------------------------------------------------
        apply_url = (
            job.get("redirect_url")
            or job.get("adref")
            or ""
        )

        if not apply_url:

            job_id = job.get("id")

            if job_id:

                apply_url = (
                    "https://www.adzuna.in/details/"
                    f"{job_id}"
                )

        # ----------------------------------------------------
        # Company
        # ----------------------------------------------------
        company = job.get("company")

        if isinstance(company, dict):
            company = company.get(
                "display_name"
            )

        # ----------------------------------------------------
        # Location
        # ----------------------------------------------------
        job_location = job.get("location")

        if isinstance(job_location, dict):
            job_location = job_location.get(
                "display_name"
            )

        # ----------------------------------------------------
        # Category
        # ----------------------------------------------------
        category = job.get("category")

        if isinstance(category, dict):
            category = category.get(
                "label"
            )

        # ----------------------------------------------------
        # Store job
        # ----------------------------------------------------
        jobs.append({

            "id": job.get("id"),

            "title": title,

            "company": company,

            "location": job_location,

            "description": description,

            "category": category,

            "contract_type": job.get(
                "contract_type"
            ),

            "contract_time": job.get(
                "contract_time"
            ),

            "salary_min": job.get(
                "salary_min"
            ),

            "salary_max": job.get(
                "salary_max"
            ),

            "salary_is_predicted": job.get(
                "salary_is_predicted"
            ),

            "created": created,

            "age_days": age_days,

            "apply_url": apply_url
        })

    return jobs


# ============================================================
# API: JOB SEARCH
# ============================================================
@app.get("/jobs/search")
def search_jobs(
    query: str = "software developer",
    location: str = "Hyderabad",
    page: int = 1,
    results_per_page: int = 10,
    max_age_days: int = 30
):
    jobs = fetch_adzuna_jobs(
        query=query,
        location=location,
        page=page,
        results_per_page=results_per_page,
        max_age_days=max_age_days
    )

    return {
        "source": "Adzuna",
        "query": query,
        "location": location,
        "max_job_age_days": max_age_days,
        "jobs_found": len(jobs),
        "jobs": jobs,
    }


# ============================================================
# JOB MATCHING ENGINE
# ============================================================

def calculate_skill_match(
    resume_skills: list[str],
    job_text: str
) -> dict:

    resume_skill_set = {
        normalize_skill(skill)
        for skill in resume_skills
    }

    job_skills = set(extract_skills(job_text))

    matching_skills = sorted(
        resume_skill_set.intersection(job_skills)
    )

    missing_skills = sorted(
        job_skills.difference(resume_skill_set)
    )

    return {
        "job_skills": sorted(job_skills),
        "matching_skills": matching_skills,
        "missing_skills": missing_skills,
    }


def calculate_title_match_score(
    resume_skills: list[str],
    job_title: str,
    job_description: str
) -> int:

    title = normalize_text(job_title)
    description = normalize_text(job_description)

    score = 0

    # Direct evidence from the job title.
    resume_skills_normalized = {
        normalize_skill(skill)
        for skill in resume_skills
    }

    for skill in resume_skills_normalized:
        if contains_skill(title, skill):
            score += 10

    # Strong role-related signals.
    role_signals = [
        "developer",
        "software",
        "programmer",
        "application",
        "web",
        "frontend",
        "backend",
        "full stack",
        "full-stack",
    ]

    if any(signal in title for signal in role_signals):
        score += 10

    # Do not let this component exceed 20.
    return min(score, 20)


def calculate_resume_evidence_score(
    resume_text: str,
    job_description: str
) -> int:

    if not resume_text or not job_description:
        return 0

    resume_lower = normalize_text(resume_text)
    job_lower = normalize_text(job_description)

    evidence_terms = [
        "project",
        "internship",
        "experience",
        "developed",
        "built",
        "implemented",
        "designed",
        "application",
        "software",
    ]

    resume_evidence = sum(
        1
        for term in evidence_terms
        if term in resume_lower
    )

    job_evidence = sum(
        1
        for term in evidence_terms
        if term in job_lower
    )

    if resume_evidence == 0 or job_evidence == 0:
        return 0

    return min(10, resume_evidence + job_evidence)


def match_resume_to_job(
    resume_text: str,
    resume_skills: list[str],
    job: dict
) -> dict:

    job_title = job.get("title") or ""
    job_description = job.get("description") or ""

    combined_job_text = (
        f"{job_title}\n{job_description}"
    )

    skill_result = calculate_skill_match(
        resume_skills,
        combined_job_text
    )

    job_skills = skill_result["job_skills"]
    matching_skills = skill_result["matching_skills"]
    missing_skills = skill_result["missing_skills"]

    # --------------------------------------------------------
    # Component 1: Skill match = 70 points
    # --------------------------------------------------------

    if job_skills:
        skill_score = (
            len(matching_skills) / len(job_skills)
        ) * 70
    else:
        skill_score = 0

    # --------------------------------------------------------
    # Component 2: Title relevance = 20 points
    # --------------------------------------------------------

    title_score = calculate_title_match_score(
        resume_skills,
        job_title,
        job_description
    )

    # --------------------------------------------------------
    # Component 3: Resume/job evidence = 10 points
    # --------------------------------------------------------

    evidence_score = calculate_resume_evidence_score(
        resume_text,
        job_description
    )

    total_score = round(
        min(
            100,
            skill_score + title_score + evidence_score
        )
    )

    # --------------------------------------------------------
    # If no recognizable job skills exist, don't pretend
    # the result is a meaningful percentage.
    # --------------------------------------------------------

    if not job_skills:
        match_percentage = None

        score_note = (
            "No recognizable skills were detected in the "
            "available job description, so a reliable "
            "skill-based percentage could not be calculated."
        )

    else:
        match_percentage = total_score

        score_note = (
            "Score combines recognizable skill overlap, "
            "job-title relevance, and available resume/job "
            "evidence. It is an estimate, not a guarantee."
        )

    additional_resume_skills = sorted(
        set(resume_skills) - set(matching_skills)
    )

    # --------------------------------------------------------
    # Match explanation
    # --------------------------------------------------------

    if match_percentage is None:
        explanation = (
            "The job description did not contain enough "
            "recognizable skills for a reliable percentage."
        )

    elif match_percentage >= 75:
        explanation = (
            "Strong recognizable skill overlap with the job."
        )

    elif match_percentage >= 50:
        explanation = (
            "Moderate skill overlap. Review the missing skills "
            "before applying."
        )

    elif match_percentage >= 25:
        explanation = (
            "Limited skill overlap. Several job requirements "
            "may be missing from the resume."
        )

    else:
        explanation = (
            "Low recognizable skill overlap with this job."
        )

    return {
        "match_percentage": match_percentage,
        "score_note": score_note,
        "explanation": explanation,
        "resume_skills": sorted(resume_skills),
        "job_required_skills": job_skills,
        "matching_skills": matching_skills,
        "missing_skills": missing_skills,
        "additional_resume_skills": additional_resume_skills,
        "score_breakdown": {
            "skill_score": round(skill_score),
            "title_relevance_score": title_score,
            "resume_job_evidence_score": evidence_score,
            "total_score": total_score,
        },
    }


# ============================================================
# API: MATCH RESUME WITH REAL JOBS
# ============================================================

@app.post("/match-jobs")
async def match_jobs(
    file: UploadFile = File(...),
    query: str = "software developer",
    location: str = "Hyderabad",
    page: int = 1,
    results_per_page: int = 10,
    max_age_days: int = 30
):

    if not file.filename:
        raise HTTPException(
            status_code=400,
            detail="No resume filename provided."
        )

    file_bytes = await file.read()

    if not file_bytes:
        raise HTTPException(
            status_code=400,
            detail="Uploaded resume is empty."
        )

    # --------------------------------------------------------
    # Parse resume
    # --------------------------------------------------------

    resume_text = extract_resume_text(
        file.filename,
        file_bytes
    )

    if len(resume_text.strip()) < 20:
        raise HTTPException(
            status_code=400,
            detail=(
                "Could not extract enough text from the resume."
            )
        )

    resume_analysis = analyze_resume_text(
        resume_text,
        file.filename
    )

    resume_skills = resume_analysis["resume_skills"]

    # --------------------------------------------------------
    # Fetch REAL jobs
    # --------------------------------------------------------

    jobs = fetch_adzuna_jobs(
        query=query,
        location=location,
        page=page,
        results_per_page=results_per_page,
        max_age_days=max_age_days
    )

    # --------------------------------------------------------
    # Match every real job
    # --------------------------------------------------------

    matched_jobs = []

    for job in jobs:

        matching = match_resume_to_job(
            resume_text=resume_text,
            resume_skills=resume_skills,
            job=job
        )

        matched_jobs.append({
            **job,
            "matching": matching
        })

    # --------------------------------------------------------
    # Sort jobs
    #
    # Jobs with an actual calculated score come first.
    # Then highest score first.
    # --------------------------------------------------------

    matched_jobs.sort(
        key=lambda item: (
            item["matching"]["match_percentage"]
            is not None,
            item["matching"]["match_percentage"]
            if item["matching"]["match_percentage"] is not None
            else -1
        ),
        reverse=True
    )

    return {
        "source": "Adzuna",

        "resume_filename": file.filename,

        "resume_characters": len(resume_text),

        "resume_skills": resume_skills,

        "resume_sections": resume_analysis[
            "resume_sections"
        ],

        "query": query,

        "location": location,

        "max_job_age_days": max_age_days,

        "jobs_found": len(matched_jobs),

        "jobs": matched_jobs,
    }


# ============================================================
# ROOT API
# ============================================================

@app.get("/")
def root():
    return {
        "project": "Real-Time Resume Job Matcher",
        "status": "running",
        "version": "2.0.0",
        "job_source": "Adzuna",

        "features": [
            "PDF resume parsing",
            "DOCX resume parsing",
            "Resume skill extraction",
            "Real job search",
            "30-day job freshness filter",
            "Job relevance filtering",
            "Explainable resume-job matching",
            "Matching skills detection",
            "Missing skills detection",
            "Score breakdown",
            "Real application links",
        ]
    }